from io import BytesIO
from datetime import date

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from auctions.models import (
    Auction,
    AuctionEvent,
    AuctionLot,
    Bid,
    Player,
    Team,
    TeamAccess,
    TeamPlayer,
    TeamWatchlist,
)
from auctions.constants import jersey_size_label, jersey_size_number

from .forms import DEFAULT_PLAYER_BASE_PRICE, PlayerForm, PublicPlayerRegistrationForm
from .models import Tournament


class PlayerBasePriceFormTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="owner", password="pass123")
        self.tournament = Tournament.objects.create(
            organizer=self.user,
            name="Base Price Cup",
            season_year=2026,
        )

    def test_new_player_form_defaults_base_price_to_10k(self):
        form = PlayerForm()
        self.assertEqual(form.fields["base_price"].initial, DEFAULT_PLAYER_BASE_PRICE)

    def test_edit_player_form_uses_database_base_price(self):
        player = Player.objects.create(
            tournament=self.tournament,
            name="Existing Player",
            base_price=25000,
            is_active=True,
        )

        form = PlayerForm(instance=player)
        self.assertEqual(form["base_price"].value(), 25000)

    def test_public_registration_defaults_base_price_to_10k(self):
        form = PublicPlayerRegistrationForm(
            data={
                "tournament": self.tournament.pk,
                "name": "Registered Player",
                "gender": Player.Gender.MALE,
                "building": "D-1",
                "flat_no": "101",
                "dob": "2000-01-01",
                "phone": "9999999999",
                "jersey_no": 7,
                "jersey_name": "REG",
                "jersey_size": Player.JerseySize.M,
            },
            files={
                "payment_screenshot": SimpleUploadedFile(
                    "payment.gif",
                    (
                        b"GIF89a\x01\x00\x01\x00\x80\x00\x00"
                        b"\x00\x00\x00\xff\xff\xff!\xf9\x04"
                        b"\x01\x00\x00\x00\x00,\x00\x00\x00"
                        b"\x00\x01\x00\x01\x00\x00\x02\x02D"
                        b"\x01\x00;"
                    ),
                    content_type="image/gif",
                )
            },
        )

        self.assertTrue(form.is_valid(), form.errors)
        player = form.save()
        self.assertEqual(player.base_price, DEFAULT_PLAYER_BASE_PRICE)
        self.assertTrue(player.is_eligible_for_auction)

    def _registration_form(self, *, name, gender, age):
        today = date.today()
        try:
            dob = today.replace(year=today.year - age)
        except ValueError:
            dob = date(today.year - age, 2, 28)

        return PublicPlayerRegistrationForm(
            data={
                "tournament": self.tournament.pk,
                "name": name,
                "gender": gender,
                "building": "D-1",
                "flat_no": "101",
                "dob": dob.isoformat(),
                "phone": "9999999999",
                "jersey_no": 7,
                "jersey_name": "REG",
                "jersey_size": Player.JerseySize.M,
            },
            files={
                "payment_screenshot": SimpleUploadedFile(
                    f"{name}.gif",
                    (
                        b"GIF89a\x01\x00\x01\x00\x80\x00\x00"
                        b"\x00\x00\x00\xff\xff\xff!\xf9\x04"
                        b"\x01\x00\x00\x00\x00,\x00\x00\x00"
                        b"\x00\x01\x00\x01\x00\x00\x02\x02D"
                        b"\x01\x00;"
                    ),
                    content_type="image/gif",
                )
            },
        )

    def test_public_registration_marks_women_as_not_auction_eligible(self):
        form = self._registration_form(name="Adult Woman", gender=Player.Gender.FEMALE, age=25)

        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.save().is_eligible_for_auction)

    def test_public_registration_marks_13_year_old_boy_as_not_auction_eligible(self):
        form = self._registration_form(name="Junior Boy", gender=Player.Gender.MALE, age=13)

        self.assertTrue(form.is_valid(), form.errors)
        self.assertFalse(form.save().is_eligible_for_auction)

    def test_jersey_size_choices_include_numeric_mapping_and_extended_sizes(self):
        choices = dict(PublicPlayerRegistrationForm().fields["jersey_size"].choices)

        self.assertEqual(choices[Player.JerseySize.XS], "XS - 34")
        self.assertEqual(choices[Player.JerseySize.M], "M - 38")
        self.assertEqual(choices[Player.JerseySize.XXL], "2XL - 44")
        self.assertEqual(choices[Player.JerseySize.XXXL], "3XL - 46")
        self.assertEqual(choices[Player.JerseySize.XXXXL], "4XL - 48")
        self.assertEqual(jersey_size_label(Player.JerseySize.XXL), "2XL")
        self.assertEqual(jersey_size_number(Player.JerseySize.XXL), 44)


class TeamOwnerDashboardTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.organizer = User.objects.create_user(username="organizer", password="pass123")
        self.owner = User.objects.create_user(username="teamowner", password="pass123")
        self.outsider = User.objects.create_user(username="outsider", password="pass123")

        self.tournament = Tournament.objects.create(
            organizer=self.organizer,
            name="Owner Portal Cup",
            season_year=2026,
        )
        self.auction = Auction.objects.create(
            tournament=self.tournament,
            name="Main Auction",
            code="SS-OWNER1",
            status=Auction.Status.LIVE,
            rules={"auction_type": "rule_based", "rule_lot1": True, "lot1_players_per_team": 2},
        )
        self.team = Team.objects.create(
            tournament=self.tournament,
            name="Sky Strikers",
            short_name="SSK",
            purse_total=2000000,
            purse_remaining=1750000,
            max_players=8,
            tagline="Own the night.",
            primary_color="#102A43",
            secondary_color="#2EC4B6",
            accent_color="#FF9F1C",
        )
        TeamAccess.objects.create(
            team=self.team,
            user=self.owner,
            role=TeamAccess.Role.OWNER,
            is_primary=True,
            is_active=True,
            linked_player=None,
        )

        owner_player = Player.objects.create(
            tournament=self.tournament,
            name="Owner On Team",
            role=Player.Role.WK,
            base_price=90000,
            is_active=True,
        )
        bought_player = Player.objects.create(
            tournament=self.tournament,
            name="Captain Cool",
            role=Player.Role.BAT,
            base_price=100000,
            is_active=True,
        )
        current_player = Player.objects.create(
            tournament=self.tournament,
            name="Live Prospect",
            role=Player.Role.AR,
            base_price=120000,
            is_active=True,
        )

        AuctionLot.objects.create(
            auction=self.auction,
            player=bought_player,
            lot_no=1,
            lot_order=1,
            status=AuctionLot.Status.SOLD,
            base_price_snapshot=100000,
            sold_to_team=self.team,
            sold_price=250000,
        )
        running_lot = AuctionLot.objects.create(
            auction=self.auction,
            player=current_player,
            lot_no=1,
            lot_order=2,
            status=AuctionLot.Status.RUNNING,
            base_price_snapshot=120000,
        )

        TeamPlayer.objects.create(
            tournament=self.tournament,
            team=self.team,
            player=owner_player,
            bought_price=80000,
            designation=TeamPlayer.Designation.OWNER,
        )

        TeamPlayer.objects.create(
            tournament=self.tournament,
            team=self.team,
            player=bought_player,
            bought_in_auction=self.auction,
            bought_price=250000,
            designation=TeamPlayer.Designation.CAPTAIN,
        )

        access = TeamAccess.objects.get(team=self.team, user=self.owner)
        access.linked_player = owner_player
        access.save(update_fields=["linked_player"])

        self.owner_player = owner_player
        self.bought_player = bought_player
        self.current_player = current_player
        self.running_lot = running_lot

    def test_owner_can_view_owner_dashboard(self):
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Sky Strikers")
        self.assertContains(response, "Captain Cool")
        self.assertContains(response, "Live Prospect")
        self.assertContains(response, "Can Bid")
        self.assertContains(response, "Browse All Players")
        self.assertContains(response, "All Players")
        self.assertContains(response, "All Teams")
        self.assertContains(
            response,
            f'{reverse("teams_dashboard", kwargs={"slug": self.tournament.slug})}?source=owner',
        )

    def test_owner_dashboard_hides_deferred_strategy_panels(self):
        self.client.force_login(self.owner)
        Bid.objects.create(
            lot=self.running_lot,
            team=self.team,
            amount=120000,
            is_valid=True,
        )

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Watchlist")
        self.assertNotContains(response, "Spending Plan")
        self.assertNotContains(response, "Demand Watch")
        self.assertNotContains(response, "Role Gaps")
        self.assertNotContains(response, "Bargain Candidates")
        self.assertNotContains(response, "Stage Mode")
        self.assertNotContains(response, "Edit Team")

    def test_owner_player_catalog_shows_statuses_and_stats_without_photos(self):
        unsold_player = Player.objects.create(
            tournament=self.tournament,
            name="Unsold Searchable Surname",
            role=Player.Role.BOWL,
            base_price=60000,
            stats={"matches": 8, "wickets": 12},
            phone="9876543210",
            is_active=True,
        )
        AuctionLot.objects.create(
            auction=self.auction,
            player=unsold_player,
            lot_no=2,
            lot_order=1,
            status=AuctionLot.Status.UNSOLD,
            base_price_snapshot=60000,
        )
        self.current_player.stats = {"matches": 10, "runs": 245}
        self.current_player.save(update_fields=["stats"])
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertContains(response, "Sold")
        self.assertContains(response, "Retained")
        self.assertContains(response, "Unsold")
        self.assertContains(response, "Remain to Auction")
        self.assertContains(response, "Unsold Searchable Surname")
        self.assertContains(response, "Matches:")
        self.assertContains(response, "Wickets:")
        self.assertContains(response, "<th>Contact</th>", html=True)
        self.assertContains(response, "9876543210")
        self.assertNotContains(response, "<th>Role</th>", html=True)
        self.assertNotContains(response, "<th>Base Price</th>", html=True)
        self.assertNotContains(response, "<th>Lot</th>", html=True)
        self.assertNotContains(response, "player.photo")

    def test_player_export_includes_contact_and_both_jersey_size_formats(self):
        player = Player.objects.create(
            tournament=self.tournament,
            name="Jersey Export Player",
            phone="9123456780",
            jersey_name="EXPORTER",
            jersey_no=27,
            jersey_size=Player.JerseySize.XXL,
            is_active=True,
        )
        self.client.force_login(self.organizer)

        response = self.client.get(
            reverse("players_export_xlsx", kwargs={"slug": self.tournament.slug})
        )

        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(response.content), read_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        self.assertEqual(
            rows[0],
            ("Name", "Contact", "Jersey Name", "Jersey No", "Jersey Size", "Jersey Size No"),
        )
        exported = next(row for row in rows[1:] if row[0] == player.name)
        self.assertEqual(
            exported,
            (player.name, "9123456780", "EXPORTER", 27, "2XL", 44),
        )

    def test_eligible_player_export_excludes_ineligible_registrations(self):
        eligible = Player.objects.create(
            tournament=self.tournament,
            name="Eligible Export Player",
            is_eligible_for_auction=True,
        )
        ineligible = Player.objects.create(
            tournament=self.tournament,
            name="All Players Export Only",
            is_eligible_for_auction=False,
        )
        self.client.force_login(self.organizer)

        all_response = self.client.get(
            reverse("players_export_xlsx", kwargs={"slug": self.tournament.slug})
        )
        eligible_response = self.client.get(
            reverse("eligible_players_export_xlsx", kwargs={"slug": self.tournament.slug})
        )

        from openpyxl import load_workbook

        all_rows = list(load_workbook(BytesIO(all_response.content), read_only=True).active.iter_rows(values_only=True))
        eligible_rows = list(
            load_workbook(BytesIO(eligible_response.content), read_only=True).active.iter_rows(values_only=True)
        )
        self.assertIn(eligible.name, {row[0] for row in all_rows[1:]})
        self.assertIn(ineligible.name, {row[0] for row in all_rows[1:]})
        self.assertIn(eligible.name, {row[0] for row in eligible_rows[1:]})
        self.assertNotIn(ineligible.name, {row[0] for row in eligible_rows[1:]})

    def test_owner_dashboard_and_lot_manager_hide_ineligible_players(self):
        ineligible = Player.objects.create(
            tournament=self.tournament,
            name="Not For Auction Player",
            is_active=True,
            is_eligible_for_auction=False,
        )
        AuctionLot.objects.create(
            auction=self.auction,
            player=ineligible,
            lot_no=4,
            lot_order=1,
            status=AuctionLot.Status.PENDING,
        )

        self.client.force_login(self.owner)
        owner_response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )
        self.assertNotContains(owner_response, ineligible.name)

        self.client.force_login(self.organizer)
        lot_response = self.client.get(reverse("lot_manager", kwargs={"slug": self.tournament.slug}))
        self.assertNotContains(lot_response, ineligible.name)

    def test_running_watchlist_player_highlights_current_lot(self):
        TeamWatchlist.objects.create(
            team=self.team,
            player=self.current_player,
            added_by=self.owner,
            priority=TeamWatchlist.Priority.HIGH,
            note="Push up to 3 lakh",
        )
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertContains(response, "Watchlist target is live")
        self.assertContains(response, "Push up to 3 lakh")
        self.assertContains(response, "watchlist-live")

    def test_started_pending_watchlist_player_highlights_current_lot(self):
        self.running_lot.status = AuctionLot.Status.PENDING
        self.running_lot.save(update_fields=["status"])
        AuctionEvent.objects.create(
            auction=self.auction,
            lot=self.running_lot,
            player=self.current_player,
            event_type=AuctionEvent.EventType.LOT_START,
            level=AuctionEvent.Level.INFO,
            message="Live Prospect is up for auction.",
        )
        TeamWatchlist.objects.create(
            team=self.team,
            player=self.current_player,
            added_by=self.owner,
            priority=TeamWatchlist.Priority.HIGH,
        )
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertContains(response, "Watchlist target is live")
        self.assertContains(response, "watchlist-live")

    def test_unstarted_pending_watchlist_player_does_not_highlight(self):
        self.running_lot.status = AuctionLot.Status.PENDING
        self.running_lot.save(update_fields=["status"])
        TeamWatchlist.objects.create(
            team=self.team,
            player=self.current_player,
            added_by=self.owner,
            priority=TeamWatchlist.Priority.HIGH,
        )
        self.client.force_login(self.owner)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertNotContains(response, "Watchlist target is live")

    def test_owner_can_add_and_remove_watchlist_player(self):
        self.client.force_login(self.owner)
        target = Player.objects.create(
            tournament=self.tournament,
            name="Watch Target",
            role=Player.Role.BOWL,
            base_price=50000,
            is_active=True,
        )

        response = self.client.post(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
            data={
                "action": "add_watchlist",
                "player_id": target.id,
                "priority": TeamWatchlist.Priority.HIGH,
                "note": "Death overs option",
            },
        )

        self.assertRedirects(
            response,
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
        )
        entry = TeamWatchlist.objects.get(team=self.team, player=target)
        self.assertEqual(entry.priority, TeamWatchlist.Priority.HIGH)
        self.assertEqual(entry.note, "Death overs option")

        response = self.client.post(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
            data={
                "action": "remove_watchlist",
                "watchlist_entry_id": entry.id,
            },
        )

        self.assertRedirects(
            response,
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
        )
        self.assertFalse(TeamWatchlist.objects.filter(team=self.team, player=target).exists())

    def test_owner_cannot_add_sold_player_to_watchlist(self):
        self.client.force_login(self.owner)

        response = self.client.post(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
            data={
                "action": "add_watchlist",
                "player_id": self.bought_player.id,
                "priority": TeamWatchlist.Priority.HIGH,
            },
            follow=True,
        )

        self.assertContains(response, "Sold or retained players cannot be added")
        self.assertFalse(
            TeamWatchlist.objects.filter(team=self.team, player=self.bought_player).exists()
        )

    def test_team_owner_can_open_all_teams_dashboard(self):
        self.client.force_login(self.owner)
        teams_url = reverse("teams_dashboard", kwargs={"slug": self.tournament.slug})

        response = self.client.get(f"{teams_url}?source=owner")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Sky Strikers")
        self.assertNotContains(
            response,
            reverse("auction_screen", kwargs={"slug": self.tournament.slug}),
        )
        self.assertNotContains(
            response,
            f'href="{reverse("tournament_detail", kwargs={"slug": self.tournament.slug})}"',
        )
        self.assertNotContains(
            response,
            reverse("teams_export_xlsx", kwargs={"slug": self.tournament.slug}),
        )

    def test_auction_source_all_teams_shows_only_export(self):
        self.client.force_login(self.organizer)
        teams_url = reverse("teams_dashboard", kwargs={"slug": self.tournament.slug})

        response = self.client.get(f"{teams_url}?source=auction")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Export Teams")
        self.assertNotContains(response, "Back to Auction")
        self.assertNotContains(
            response,
            f'href="{reverse("tournament_detail", kwargs={"slug": self.tournament.slug})}"',
        )

    def test_outsider_cannot_open_all_teams_dashboard(self):
        self.client.force_login(self.outsider)

        response = self.client.get(
            reverse("teams_dashboard", kwargs={"slug": self.tournament.slug})
        )

        self.assertRedirects(response, reverse("tournament_list"))

    def test_auction_teams_link_opens_in_new_tab(self):
        self.client.force_login(self.organizer)

        response = self.client.get(
            reverse("auction_screen", kwargs={"slug": self.tournament.slug})
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(
            response,
            f'href="{reverse("teams_dashboard", kwargs={"slug": self.tournament.slug})}?source=auction"',
        )
        self.assertContains(response, 'target="_blank"')

    def test_tournament_list_redirects_single_owner_to_dashboard(self):
        self.client.force_login(self.owner)

        response = self.client.get(reverse("tournament_list"))

        self.assertRedirects(
            response,
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
        )

    def test_owner_dashboard_blocks_unlinked_user(self):
        self.client.force_login(self.outsider)

        response = self.client.get(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id})
        )

        self.assertRedirects(response, reverse("tournament_list"))

    def test_owner_can_update_other_player_designation_to_captain(self):
        self.client.force_login(self.owner)
        team_player = TeamPlayer.objects.get(team=self.team, player=self.bought_player)

        response = self.client.post(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
            data={
                "action": "update_owner_squad",
                "team_player_id": team_player.id,
                "designation": TeamPlayer.Designation.CAPTAIN,
            },
        )

        self.assertRedirects(
            response,
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
        )
        team_player.refresh_from_db()
        self.assertEqual(team_player.designation, TeamPlayer.Designation.CAPTAIN)

    def test_owner_cannot_change_own_designation(self):
        self.client.force_login(self.owner)
        own_entry = TeamPlayer.objects.get(team=self.team, player=self.owner_player)

        response = self.client.post(
            reverse("owner_dashboard", kwargs={"slug": self.tournament.slug, "team_id": self.team.id}),
            data={
                "action": "update_owner_squad",
                "team_player_id": own_entry.id,
                "designation": TeamPlayer.Designation.CAPTAIN,
            },
            follow=True,
        )

        own_entry.refresh_from_db()

        self.assertEqual(own_entry.designation, TeamPlayer.Designation.OWNER)
        self.assertContains(response, "cannot change your own designation")
