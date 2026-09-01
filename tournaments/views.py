import json

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import IntegrityError, transaction
from django.db.models import Avg, Count, Max, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.template.loader import render_to_string
from django.utils.crypto import get_random_string
from django.utils import timezone
from django.views.decorators.http import require_POST

from decimal import Decimal, InvalidOperation

from auctions.models import Auction, Team, TeamAccess, TeamWatchlist, TeamPlayer, Player, AuctionLot, Bid
from auctions.rules import (
    AUCTION_TYPE_OPEN,
    AUCTION_TYPE_RULE_BASED,
    PRICE_CAP_VALUES,
    RULE_LOT1,
    RULE_LOT1_AND_LOT2,
    RULE_PRICE_CAPS,
    default_auction_rules,
    normalize_auction_rules,
)
from auctions.services import (
    generate_lots_from_active_players,
    generate_random_lots_from_active_players,
    resequence_lot_orders,
)

from .forms import TeamAccessForm, TeamCreateForm, TournamentCreateForm, PlayerForm
from .models import Tournament
from .permissions import can_access_team_dashboard, can_manage_tournament, user_team_access_qs


def _ensure_auction(tournament: Tournament) -> Auction:
    auction = getattr(tournament, "auction", None)
    if auction:
        return auction

    season = tournament.season_year or ""
    prefix = f"SS-{season}-" if season else "SS-"

    for _ in range(20):
        code = f"{prefix}{get_random_string(6).upper()}"
        try:
            return Auction.objects.create(
                tournament=tournament,
                name="Main Auction",
                code=code,
                status=Auction.Status.DRAFT,
            )
        except IntegrityError:
            continue

    raise RuntimeError("Failed to generate a unique auction code")


def _filter_tournament_qs(qs, user):
    # Accept either a QuerySet or a Model class (get_object_or_404 supports both).
    if hasattr(qs, "_default_manager"):
        qs = qs._default_manager.all()

    if user.is_superuser:
        return qs
    return qs.filter(organizer=user)


def _is_ajax(request) -> bool:
    return request.headers.get("x-requested-with") == "XMLHttpRequest"


def _auction_has_activity(auction: Auction | None) -> bool:
    if not auction:
        return False

    if Bid.objects.filter(lot__auction=auction, is_valid=True).exists():
        return True

    return AuctionLot.objects.filter(
        auction=auction,
        status__in=[
            AuctionLot.Status.RUNNING,
            AuctionLot.Status.SOLD,
            AuctionLot.Status.UNSOLD,
        ],
    ).exists()


def _available_lot_count_for_rules(auction: Auction, lot_no: int) -> int:
    assigned_player_ids = (
        TeamPlayer.objects.filter(tournament=auction.tournament)
        .exclude(status=TeamPlayer.Status.RELEASED)
        .values_list("player_id", flat=True)
    )
    return (
        AuctionLot.objects.filter(auction=auction, lot_no=lot_no)
        .exclude(status=AuctionLot.Status.WITHDRAWN)
        .exclude(player_id__in=assigned_player_ids)
        .count()
    )


def _parse_positive_int(raw_value):
    try:
        value = int(raw_value)
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _build_rule_validation_error(auction: Auction, lot_no: int, players_per_team: int) -> str | None:
    active_teams = Team.objects.filter(tournament=auction.tournament, is_active=True).count()
    if active_teams < 1:
        return "Add at least one active team before starting the auction."

    lot_count = _available_lot_count_for_rules(auction, lot_no)
    if lot_count < 1:
        return f"Lot {lot_no} has no auction players available for this rule."

    expected_total = active_teams * players_per_team
    if lot_count != expected_total:
        if lot_count % active_teams == 0:
            exact_value = lot_count // active_teams
            return (
                f"Lot {lot_no} has {lot_count} available players and {active_teams} active teams. "
                f"Use {exact_value} player(s) per team for an exact division."
            )
        return (
            f"Lot {lot_no} has {lot_count} available players and {active_teams} active teams, "
            "so it cannot be divided equally across teams right now."
        )

    return None


def _owner_access_cards(user):
    entries = []
    for access in user_team_access_qs(user).order_by(
            "team__tournament__name",
            "team__name",
            "-is_primary",
            "role",
    ):
        entries.append(
            {
                "access": access,
                "team": access.team,
                "tournament": access.team.tournament,
            }
        )
    return entries


def _restricted_lot_no_for_owner_view(auction: Auction) -> int | None:
    rules = normalize_auction_rules(auction)
    if not rules["is_rule_based"]:
        return None

    lot_sequence = []
    if rules["rule_lot1_and_lot2_enabled"]:
        lot_sequence = [1, 2]
    elif rules["rule_lot1_enabled"]:
        lot_sequence = [1]

    assigned_player_ids = (
        TeamPlayer.objects.filter(tournament=auction.tournament)
        .exclude(status=TeamPlayer.Status.RELEASED)
        .values_list("player_id", flat=True)
    )

    for lot_no in lot_sequence:
        has_open_or_unsold = AuctionLot.objects.filter(
            auction=auction,
            lot_no=lot_no,
            status__in=[
                AuctionLot.Status.PENDING,
                AuctionLot.Status.RUNNING,
                AuctionLot.Status.UNSOLD,
            ],
        ).exclude(player_id__in=assigned_player_ids).exists()
        if has_open_or_unsold:
            return lot_no

    return None


def _current_lot_for_owner_view(auction: Auction | None):
    if not auction:
        return None

    restricted_lot_no = _restricted_lot_no_for_owner_view(auction)
    lot_scope = AuctionLot.objects.filter(auction=auction)
    if restricted_lot_no:
        lot_scope = lot_scope.filter(lot_no=restricted_lot_no)

    current_lot = (
        lot_scope.filter(status=AuctionLot.Status.RUNNING)
        .select_related("player", "sold_to_team")
        .order_by("lot_no", "lot_order", "id")
        .first()
    )
    if current_lot:
        return current_lot

    return (
        lot_scope.filter(status=AuctionLot.Status.PENDING)
        .select_related("player", "sold_to_team")
        .order_by("lot_no", "lot_order", "id")
        .first()
    )


def _lot_count_for_team(team: Team, auction: Auction, lot_no: int) -> int:
    return (
        TeamPlayer.objects.filter(
            tournament=auction.tournament,
            team=team,
            player__lots__auction=auction,
            player__lots__lot_no=lot_no,
        )
        .exclude(status=TeamPlayer.Status.RELEASED)
        .distinct()
        .count()
    )


def _price_cap_count_for_team(team: Team, auction: Auction, base_price: Decimal) -> int:
    return (
        TeamPlayer.objects.filter(
            tournament=auction.tournament,
            team=team,
            player__base_price=base_price,
        )
        .exclude(status=TeamPlayer.Status.RELEASED)
        .count()
    )


def _team_rule_error(team: Team, auction: Auction, lot: AuctionLot, rules: dict) -> str | None:
    lot_limit = None
    if lot.lot_no == 1 and (rules["rule_lot1_enabled"] or rules["rule_lot1_and_lot2_enabled"]):
        lot_limit = rules["lot1_players_per_team"]
    elif lot.lot_no == 2 and rules["rule_lot1_and_lot2_enabled"]:
        lot_limit = rules["lot2_players_per_team"]

    if lot_limit:
        lot_count = _lot_count_for_team(team, auction, lot.lot_no)
        if lot_count >= lot_limit:
            return f"Blocked by Lot {lot.lot_no} rule. Limit reached for this lot."

    if rules["rule_price_caps_enabled"] and int(lot.base_price_snapshot or 0) in PRICE_CAP_VALUES:
        base_price = Decimal(lot.base_price_snapshot or 0)
        if _price_cap_count_for_team(team, auction, base_price) >= 1:
            return f"Blocked by price cap. Team already has a {int(base_price):,} bracket player."

    return None


def _next_bid_amount(lot: AuctionLot) -> Decimal:
    top_amount = (
        Bid.objects.filter(lot=lot, is_valid=True)
        .order_by("-amount", "-bid_time")
        .values_list("amount", flat=True)
        .first()
    )
    if top_amount is None:
        return Decimal(lot.base_price_snapshot or 0)

    current = Decimal(top_amount)
    if current < Decimal("100000"):
        return current + Decimal("10000")
    if current < Decimal("400000"):
        return current + Decimal("20000")
    return current + Decimal("50000")


def _current_highest_bid(lot: AuctionLot):
    return (
        Bid.objects.filter(lot=lot, is_valid=True)
        .select_related("team")
        .order_by("-amount", "-bid_time")
        .first()
    )


def _appearance_no(lot: AuctionLot) -> int:
    before = AuctionLot.objects.filter(auction=lot.auction).filter(
        Q(lot_no__lt=lot.lot_no)
        | Q(lot_no=lot.lot_no, lot_order__lt=lot.lot_order)
        | Q(lot_no=lot.lot_no, lot_order=lot.lot_order, id__lt=lot.id)
    )
    return before.count() + 1


def _owner_bid_status(team: Team, current_lot: AuctionLot | None, players_count: int):
    if not current_lot or current_lot.status not in [AuctionLot.Status.PENDING, AuctionLot.Status.RUNNING]:
        return {
            "state": "idle",
            "label": "Waiting",
            "detail": "No live or upcoming lot is open right now.",
            "next_amount": None,
            "highest_bid": None,
            "highest_team": None,
            "is_leading": False,
        }

    highest_bid = _current_highest_bid(current_lot)
    highest_team = highest_bid.team if highest_bid else None
    next_amount = _next_bid_amount(current_lot)
    rules = normalize_auction_rules(current_lot.auction)

    rule_error = _team_rule_error(team, current_lot.auction, current_lot, rules)
    if rule_error:
        return {
            "state": "blocked",
            "label": "Blocked",
            "detail": rule_error,
            "next_amount": next_amount,
            "highest_bid": highest_bid,
            "highest_team": highest_team,
            "is_leading": highest_team and highest_team.id == team.id,
        }

    if team.max_players and players_count >= team.max_players:
        return {
            "state": "blocked",
            "label": "Blocked",
            "detail": "Blocked by squad limit. Team has already reached max players.",
            "next_amount": next_amount,
            "highest_bid": highest_bid,
            "highest_team": highest_team,
            "is_leading": highest_team and highest_team.id == team.id,
        }

    if (team.purse_total or 0) > 0:
        purse_remaining = Decimal(team.purse_remaining or 0)

        if team.max_players:
            remaining_after_purchase = max(team.max_players - (players_count + 1), 0)
            reserve_required = Decimal("10000") * Decimal(remaining_after_purchase)
            required_total = next_amount + reserve_required
            if purse_remaining < required_total:
                return {
                    "state": "blocked",
                    "label": "Blocked",
                    "detail": (
                        f"Blocked by purse reserve. Needs at least {required_total:,.0f} "
                        f"to cover the next bid and remaining slots."
                    ),
                    "next_amount": next_amount,
                    "highest_bid": highest_bid,
                    "highest_team": highest_team,
                    "is_leading": highest_team and highest_team.id == team.id,
                }

        if purse_remaining < next_amount:
            return {
                "state": "blocked",
                "label": "Blocked",
                "detail": "Blocked by purse. Team does not have enough balance for the next bid.",
                "next_amount": next_amount,
                "highest_bid": highest_bid,
                "highest_team": highest_team,
                "is_leading": highest_team and highest_team.id == team.id,
            }

    is_leading = highest_team and highest_team.id == team.id
    return {
        "state": "leading" if is_leading else "eligible",
        "label": "Leading" if is_leading else "Can Bid",
        "detail": (
            "Your team currently has the highest bid on this player."
            if is_leading
            else "This team is eligible to bid on the current lot."
        ),
        "next_amount": next_amount,
        "highest_bid": highest_bid,
        "highest_team": highest_team,
        "is_leading": is_leading,
    }


def _owner_access_for_team(user, team: Team):
    if can_manage_tournament(user, team.tournament):
        return None
    return user_team_access_qs(user).filter(team=team).select_related("linked_player").first()


def _can_edit_owner_squad(user, team: Team, owner_access) -> bool:
    if can_manage_tournament(user, team.tournament):
        return True
    if not owner_access:
        return False
    return owner_access.role in {TeamAccess.Role.OWNER, TeamAccess.Role.CO_OWNER}


def _can_edit_owner_strategy(user, team: Team, owner_access) -> bool:
    if can_manage_tournament(user, team.tournament):
        return True
    if not owner_access:
        return False
    return owner_access.role in {TeamAccess.Role.OWNER, TeamAccess.Role.CO_OWNER, TeamAccess.Role.ANALYST}


def _role_gap_indicators(squad: list[TeamPlayer]) -> list[dict]:
    role_targets = {
        Player.Role.BAT: 3,
        Player.Role.BOWL: 3,
        Player.Role.AR: 2,
        Player.Role.WK: 1,
    }
    counts = {value: 0 for value, _label in Player.Role.choices}
    for entry in squad:
        if entry.player.role in counts:
            counts[entry.player.role] += 1

    indicators = []
    for value, label in Player.Role.choices:
        target = role_targets.get(value, 0)
        count = counts.get(value, 0)
        gap = max(target - count, 0)
        if gap >= 2:
            status = "urgent"
        elif gap == 1:
            status = "open"
        else:
            status = "filled"
        indicators.append(
            {
                "value": value,
                "label": label,
                "count": count,
                "target": target,
                "gap": gap,
                "status": status,
                "width": min(100, int((count / target) * 100)) if target else 100,
            }
        )
    return indicators


def _player_lot_map(auction: Auction | None) -> dict[int, AuctionLot]:
    if not auction:
        return {}
    return {
        lot.player_id: lot
        for lot in AuctionLot.objects.filter(auction=auction)
        .select_related("player")
        .order_by("lot_no", "lot_order", "id")
    }


def _build_owner_dashboard_context(user, team: Team) -> dict:
    tournament = team.tournament
    auction = getattr(tournament, "auction", None)
    owner_access = _owner_access_for_team(user, team)
    can_edit_squad = _can_edit_owner_squad(user, team, owner_access)
    can_edit_strategy = _can_edit_owner_strategy(user, team, owner_access)
    self_locked_player_id = owner_access.linked_player_id if owner_access else None

    squad_qs = (
        TeamPlayer.objects.filter(tournament=tournament, team=team)
        .exclude(status=TeamPlayer.Status.RELEASED)
        .select_related("player")
        .order_by("-designation", "-bought_price", "player__name")
    )
    squad = list(squad_qs)

    summary = squad_qs.aggregate(
        spent=Sum("bought_price"),
        average_buy=Avg("bought_price"),
        players_count=Count("id"),
    )
    players_count = int(summary["players_count"] or 0)
    spent = Decimal(summary["spent"] or 0)
    slots_remaining = max((team.max_players or 0) - players_count, 0) if team.max_players else None

    priced_entries = [tp for tp in squad if tp.bought_price is not None]
    highest_buy_entry = max(priced_entries, key=lambda tp: tp.bought_price, default=None)
    lowest_buy_entry = min(priced_entries, key=lambda tp: tp.bought_price, default=None)
    captain_entry = next((tp for tp in squad if tp.designation == TeamPlayer.Designation.CAPTAIN), None)
    marquee_entry = next((tp for tp in squad if tp.designation == TeamPlayer.Designation.MARQUEE), None)

    lot_breakdown = []
    if auction:
        rules = normalize_auction_rules(auction)
        sold_lots = list(
            AuctionLot.objects.filter(
                auction=auction,
                sold_to_team=team,
                status=AuctionLot.Status.SOLD,
            )
            .select_related("player")
            .order_by("lot_no", "lot_order", "player__name")
        )
        lot_breakdown_map = {}
        for lot in sold_lots:
            item = lot_breakdown_map.setdefault(
                lot.lot_no,
                {
                    "lot_no": lot.lot_no,
                    "count": 0,
                    "limit": None,
                    "players": [],
                },
            )
            item["count"] += 1
            item["players"].append(lot.player.name)

        lot_numbers = set(lot_breakdown_map.keys())
        if rules["rule_lot1_enabled"] or rules["rule_lot1_and_lot2_enabled"]:
            lot_numbers.add(1)
        if rules["rule_lot1_and_lot2_enabled"]:
            lot_numbers.add(2)

        for lot_no in sorted(lot_numbers):
            item = lot_breakdown_map.get(
                lot_no,
                {
                    "lot_no": lot_no,
                    "count": 0,
                    "limit": None,
                    "players": [],
                },
            )
            if lot_no == 1 and (rules["rule_lot1_enabled"] or rules["rule_lot1_and_lot2_enabled"]):
                item["limit"] = rules["lot1_players_per_team"]
            elif lot_no == 2 and rules["rule_lot1_and_lot2_enabled"]:
                item["limit"] = rules["lot2_players_per_team"]
            lot_breakdown.append(item)

    current_lot = _current_lot_for_owner_view(auction)
    bid_status = _owner_bid_status(team, current_lot, players_count)
    current_lot_payload = None
    if current_lot:
        current_lot_payload = {
            "lot": current_lot,
            "appearance_no": _appearance_no(current_lot),
            "highest_bid": bid_status["highest_bid"],
            "highest_team": bid_status["highest_team"],
            "next_amount": bid_status["next_amount"],
        }

    recent_team_events = []
    if auction:
        recent_team_events = list(
            auction.events.filter(Q(team=team) | Q(lot__sold_to_team=team))
            .select_related("player", "team", "lot")
            .order_by("-id")[:8]
        )

    active_assignment_ids = set(
        TeamPlayer.objects.filter(tournament=tournament)
        .exclude(status=TeamPlayer.Status.RELEASED)
        .values_list("player_id", flat=True)
    )
    lot_by_player_id = _player_lot_map(auction)
    watchlist_entries = list(
        TeamWatchlist.objects.filter(team=team)
        .select_related("player", "added_by")
        .order_by("priority", "player__base_price", "player__name")
    )
    watched_player_ids = {entry.player_id for entry in watchlist_entries}
    watchlist_candidates = list(
        Player.objects.filter(tournament=tournament, is_active=True)
        .exclude(id__in=active_assignment_ids)
        .exclude(id__in=watched_player_ids)
        .order_by("base_price", "name")[:40]
    )

    role_gap_indicators = _role_gap_indicators(squad)

    purse_total = Decimal(team.purse_total or 0)
    purse_remaining = Decimal(team.purse_remaining or 0)
    purse_used_percent = int((spent / purse_total) * Decimal("100")) if purse_total else 0
    effective_slots_remaining = slots_remaining if slots_remaining is not None else 0
    average_per_open_slot = (
        purse_remaining / Decimal(effective_slots_remaining)
        if effective_slots_remaining > 0
        else None
    )
    reserve_after_next_buy = (
        Decimal("10000") * Decimal(max(effective_slots_remaining - 1, 0))
        if effective_slots_remaining
        else Decimal("0")
    )
    aggressive_next_bid = max(purse_remaining - reserve_after_next_buy, Decimal("0"))
    spending_insights = {
        "purse_used_percent": min(100, purse_used_percent),
        "average_per_open_slot": average_per_open_slot,
        "reserve_after_next_buy": reserve_after_next_buy,
        "aggressive_next_bid": aggressive_next_bid,
        "squad_floor_needed": Decimal("10000") * Decimal(effective_slots_remaining),
    }

    available_players_qs = Player.objects.filter(tournament=tournament, is_active=True).exclude(id__in=active_assignment_ids)
    if auction:
        available_players_qs = available_players_qs.exclude(
            lots__auction=auction,
            lots__status__in=[AuctionLot.Status.SOLD, AuctionLot.Status.WITHDRAWN],
        )

    value_limit = average_per_open_slot or purse_remaining
    bargain_candidates = list(
        available_players_qs.filter(base_price__lte=value_limit)
        .annotate(
            bid_count=Count("lots__bids", filter=Q(lots__bids__is_valid=True)),
            highest_market_bid=Max("lots__bids__amount", filter=Q(lots__bids__is_valid=True)),
        )
        .order_by("base_price", "-bid_count", "name")[:6]
    )
    demand_players = list(
        Player.objects.filter(tournament=tournament, is_active=True)
        .exclude(id__in=active_assignment_ids)
        .annotate(
            bid_count=Count("lots__bids", filter=Q(lots__bids__is_valid=True)),
            highest_market_bid=Max("lots__bids__amount", filter=Q(lots__bids__is_valid=True)),
        )
        .filter(bid_count__gt=0)
        .order_by("-bid_count", "-highest_market_bid", "name")[:6]
    )

    for player in watchlist_entries:
        player.auction_lot = lot_by_player_id.get(player.player_id)
    for player in watchlist_candidates:
        player.auction_lot = lot_by_player_id.get(player.id)
    for player in bargain_candidates:
        player.auction_lot = lot_by_player_id.get(player.id)
    for player in demand_players:
        player.auction_lot = lot_by_player_id.get(player.id)

    return {
        "tournament": tournament,
        "team": team,
        "auction": auction,
        "squad": squad,
        "players_count": players_count,
        "spent": spent,
        "slots_remaining": slots_remaining,
        "average_buy": summary["average_buy"],
        "highest_buy_entry": highest_buy_entry,
        "lowest_buy_entry": lowest_buy_entry,
        "captain_entry": captain_entry,
        "marquee_entry": marquee_entry,
        "lot_breakdown": lot_breakdown,
        "current_lot_payload": current_lot_payload,
        "bid_status": bid_status,
        "recent_team_events": recent_team_events,
        "auction_rules_config": normalize_auction_rules(auction) if auction else default_auction_rules(),
        "can_manage": can_manage_tournament(user, tournament),
        "can_edit_squad": can_edit_squad,
        "can_edit_strategy": can_edit_strategy,
        "self_locked_player_id": self_locked_player_id,
        "designation_choices": TeamPlayer.Designation.choices,
        "watchlist_priority_choices": TeamWatchlist.Priority.choices,
        "watchlist_entries": watchlist_entries,
        "watchlist_candidates": watchlist_candidates,
        "role_gap_indicators": role_gap_indicators,
        "spending_insights": spending_insights,
        "bargain_candidates": bargain_candidates,
        "demand_players": demand_players,
        "owner_access": owner_access,
    }


@login_required
def tournament_list(request):
    qs = Tournament.objects.select_related("organizer").order_by("-created_at")
    tournaments = _filter_tournament_qs(qs, request.user)
    owner_access_cards = _owner_access_cards(request.user)

    if request.method == "POST":
        form = TournamentCreateForm(request.POST)
        if form.is_valid():
            tournament = form.save(commit=False)
            tournament.organizer = request.user
            tournament.save()
            _ensure_auction(tournament)
            messages.success(request, "Tournament created.")
            return redirect("tournament_detail", slug=tournament.slug)
    else:
        form = TournamentCreateForm()

    if request.method == "GET" and not tournaments.exists() and len(owner_access_cards) == 1:
        only_access = owner_access_cards[0]
        return redirect(
            "owner_dashboard",
            slug=only_access["tournament"].slug,
            team_id=only_access["team"].id,
        )

    return render(
        request,
        "tournaments/tournament_list.html",
        {
            "tournaments": tournaments,
            "create_form": form,
            "owner_access_cards": owner_access_cards,
        },
    )


@login_required
def tournament_create(request):
    if request.method == "POST":
        form = TournamentCreateForm(request.POST)
        if form.is_valid():
            tournament = form.save(commit=False)
            tournament.organizer = request.user
            tournament.save()

            _ensure_auction(tournament)

            messages.success(request, "Tournament created.")
            return redirect("tournament_detail", slug=tournament.slug)
    else:
        form = TournamentCreateForm()

    return render(request, "tournaments/tournament_form.html", {"form": form})


@login_required
def tournament_detail(request, slug: str):
    tournament = get_object_or_404(
        _filter_tournament_qs(
            Tournament.objects.select_related("auction").prefetch_related("teams", "players"),
            request.user,
        ),
        slug=slug,
    )
    teams = tournament.teams.order_by("name")
    players = tournament.players.order_by("name")

    sold_player_ids = set()
    if tournament.auction:
        sold_player_ids = set(
            AuctionLot.objects.filter(
                auction=tournament.auction,
                status=AuctionLot.Status.SOLD,
            ).values_list("player_id", flat=True)
        )

    assigned_player_ids = set(
        TeamPlayer.objects.filter(tournament=tournament).values_list("player_id", flat=True)
    )
    auction = getattr(tournament, "auction", None)
    auction_rules_config = normalize_auction_rules(auction) if auction else default_auction_rules()
    auction_setup_locked = _auction_has_activity(auction)

    context = {
        "tournament": tournament,
        "teams": teams,
        "players": players,
        "sold_player_ids": sold_player_ids,
        "assigned_player_ids": assigned_player_ids,
        "total_players": players.count(),
        "active_players": players.filter(is_active=True).count(),
        "total_teams": teams.count(),
        "can_manage": can_manage_tournament(request.user, tournament),
        "player_form": PlayerForm(),
        "team_form": TeamCreateForm(),
        "auction_rules_config": auction_rules_config,
        "auction_setup_locked": auction_setup_locked,
    }
    return render(request, "tournaments/tournament_detail.html", context)


@login_required
@require_POST
def setup_auction(request, slug: str):
    tournament = get_object_or_404(
        _filter_tournament_qs(Tournament.objects.select_related("auction"), request.user),
        slug=slug,
    )

    if not can_manage_tournament(request.user, tournament):
        messages.error(request, "Forbidden.")
        return redirect("tournament_detail", slug=tournament.slug)

    auction = _ensure_auction(tournament)
    if _auction_has_activity(auction):
        messages.error(request, "Auction rules can't be changed after auction activity has started.")
        return redirect("auction_screen", slug=tournament.slug)

    auction_type = (request.POST.get("auction_type") or AUCTION_TYPE_OPEN).strip().lower()
    if auction_type not in {AUCTION_TYPE_OPEN, AUCTION_TYPE_RULE_BASED}:
        auction_type = AUCTION_TYPE_OPEN

    if auction_type == AUCTION_TYPE_OPEN:
        auction.rules = {"auction_type": AUCTION_TYPE_OPEN}
        auction.save(update_fields=["rules", "updated_at"])
        messages.success(request, "Auction is set to open auction.")
        return redirect("auction_screen", slug=tournament.slug)

    rule_lot1 = request.POST.get("rule_lot1") == "1"
    rule_lot1_and_lot2 = request.POST.get("rule_lot1_and_lot2") == "1"
    rule_price_caps = request.POST.get("rule_price_caps") == "1"

    if rule_lot1 and rule_lot1_and_lot2:
        messages.error(request, "Select either Rule 1 or Rule 2, not both.")
        return redirect("tournament_detail", slug=tournament.slug)

    if not any([rule_lot1, rule_lot1_and_lot2, rule_price_caps]):
        messages.error(request, "Select at least one rule for a rule-based auction.")
        return redirect("tournament_detail", slug=tournament.slug)

    lot1_players_per_team = None
    lot2_players_per_team = None

    if rule_lot1:
        lot1_players_per_team = _parse_positive_int(request.POST.get("rule1_lot1_players_per_team"))
        if not lot1_players_per_team:
            messages.error(request, "Enter a valid Lot 1 players-per-team value for Rule 1.")
            return redirect("tournament_detail", slug=tournament.slug)

        error = _build_rule_validation_error(auction, 1, lot1_players_per_team)
        if error:
            messages.error(request, error)
            return redirect("tournament_detail", slug=tournament.slug)

    if rule_lot1_and_lot2:
        lot1_players_per_team = _parse_positive_int(request.POST.get("rule2_lot1_players_per_team"))
        lot2_players_per_team = _parse_positive_int(request.POST.get("rule2_lot2_players_per_team"))

        if not lot1_players_per_team or not lot2_players_per_team:
            messages.error(request, "Enter valid Lot 1 and Lot 2 players-per-team values for Rule 2.")
            return redirect("tournament_detail", slug=tournament.slug)

        for lot_no, per_team in ((1, lot1_players_per_team), (2, lot2_players_per_team)):
            error = _build_rule_validation_error(auction, lot_no, per_team)
            if error:
                messages.error(request, error)
                return redirect("tournament_detail", slug=tournament.slug)

    auction.rules = {
        "auction_type": AUCTION_TYPE_RULE_BASED,
        RULE_LOT1: rule_lot1,
        RULE_LOT1_AND_LOT2: rule_lot1_and_lot2,
        RULE_PRICE_CAPS: rule_price_caps,
        "lot1_players_per_team": lot1_players_per_team,
        "lot2_players_per_team": lot2_players_per_team,
    }
    auction.save(update_fields=["rules", "updated_at"])

    messages.success(request, "Auction rules saved.")
    return redirect("auction_screen", slug=tournament.slug)


@login_required
@require_POST
def generate_lots(request, slug: str):
    tournament = get_object_or_404(
        _filter_tournament_qs(Tournament.objects.select_related("auction"), request.user),
        slug=slug,
    )

    auction = _ensure_auction(tournament)
    created, skipped = generate_lots_from_active_players(auction)

    messages.success(
        request,
        f"✅ Lots generated. Created: {created}. Skipped (already existed): {skipped}.",
    )
    return redirect("tournament_detail", slug=tournament.slug)


@login_required
@require_POST
def generate_random_lots(request, slug: str):
    tournament = get_object_or_404(
        _filter_tournament_qs(Tournament.objects.select_related("auction"), request.user),
        slug=slug,
    )

    try:
        max_lots = int(request.POST.get("max_lots") or 0)
    except (TypeError, ValueError):
        max_lots = 0

    if max_lots < 2 or max_lots > 5:
        messages.error(request, "Please choose lots between 2 and 5.")
        return redirect("tournament_detail", slug=tournament.slug)

    auction = _ensure_auction(tournament)
    created, skipped = generate_random_lots_from_active_players(auction, max_lots=max_lots, reset=True)
    lot_count = AuctionLot.objects.filter(auction=auction).values("lot_no").distinct().count()

    messages.success(
        request,
        f"✅ Random lots generated (reset). Lots: {lot_count}. Created: {created}. Skipped (already existed): {skipped}.",
    )
    return redirect("tournament_detail", slug=tournament.slug)


@login_required
def team_create(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    if request.method == "POST":
        form = TeamCreateForm(request.POST, request.FILES)
        if form.is_valid():
            team = form.save(commit=False)
            team.tournament = tournament

            if not team.short_name:
                team.short_name = "".join([w[0] for w in team.name.split()]).upper()[:8]

            # Remaining defaults to total (handled in the form), but keep a final safety net.
            if team.purse_remaining is None:
                team.purse_remaining = team.purse_total or 0

            team.save()

            if _is_ajax(request):
                team_html = render_to_string(
                    "tournaments/partials/team_card.html",
                    {"team": team, "tournament": tournament},
                    request=request,
                )
                players_qs = tournament.players.all()
                teams_qs = tournament.teams.all()
                return JsonResponse(
                    {
                        "ok": True,
                        "team": {"id": team.id, "name": team.name},
                        "team_html": team_html,
                        "counts": {
                            "total_teams": teams_qs.count(),
                            "total_players": players_qs.count(),
                            "active_players": players_qs.filter(is_active=True).count(),
                        },
                    }
                )

            messages.success(request, "Team created.")
            return redirect("tournament_detail", slug=tournament.slug)

        if _is_ajax(request):
            return JsonResponse(
                {
                    "ok": False,
                    "error": "Could not create team. Please check the inputs.",
                    "errors": form.errors.get_json_data(),
                },
                status=400,
            )
    else:
        form = TeamCreateForm()

    return render(
        request,
        "tournaments/team_form.html",
        {"tournament": tournament, "form": form},
    )


@login_required
def team_edit(request, slug: str, team_id: int):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    if not can_manage_tournament(request.user, tournament):
        messages.error(request, "Forbidden.")
        return redirect("tournament_detail", slug=tournament.slug)

    team = get_object_or_404(Team, pk=team_id, tournament=tournament)

    action = (request.POST.get("action") or "").strip() if request.method == "POST" else ""

    # Team details save
    if request.method == "POST" and action == "save_team":
        form = TeamCreateForm(request.POST, request.FILES, instance=team)
        if form.is_valid():
            form.save()
            messages.success(request, "Team updated.")
            return redirect("team_edit", slug=tournament.slug, team_id=team.id)
        messages.error(request, "Please correct the errors below and try again.")
    else:
        form = TeamCreateForm(instance=team)

    access_form = TeamAccessForm(tournament=tournament)

    # Update designations
    if request.method == "POST" and action == "update_designations":
        squad_qs = TeamPlayer.objects.filter(tournament=tournament, team=team)
        updated = 0
        with transaction.atomic():
            for tp in squad_qs:
                key = f"designation_{tp.id}"
                new_val = (request.POST.get(key) or "").strip()
                allowed = {c[0] for c in TeamPlayer.Designation.choices}
                if new_val and new_val in allowed and tp.designation != new_val:
                    tp.designation = new_val
                    tp.save(update_fields=["designation", "updated_at"])
                    updated += 1
        messages.success(request, f"Updated designations for {updated} player(s).")
        return redirect("team_edit", slug=tournament.slug, team_id=team.id)

    if request.method == "POST" and action == "add_team_access":
        access_form = TeamAccessForm(request.POST, tournament=tournament)
        if access_form.is_valid():
            team_access = access_form.save(commit=False)
            team_access.team = team
            try:
                with transaction.atomic():
                    if team_access.is_primary:
                        TeamAccess.objects.filter(team=team, is_primary=True).update(is_primary=False)
                    team_access.save()
            except IntegrityError:
                access_form.add_error("user", "This user already has access to the team.")
            else:
                messages.success(request, f"Granted {team_access.get_role_display()} access to {team_access.user}.")
                return redirect("team_edit", slug=tournament.slug, team_id=team.id)
        messages.error(request, "Could not add team access. Please check the inputs.")

    if request.method == "POST" and action == "remove_team_access":
        try:
            access_id = int(request.POST.get("access_id") or 0)
        except (TypeError, ValueError):
            access_id = 0

        team_access = get_object_or_404(TeamAccess, pk=access_id, team=team)
        removed_user = str(team_access.user)
        team_access.delete()
        messages.success(request, f"Removed team access for {removed_user}.")
        return redirect("team_edit", slug=tournament.slug, team_id=team.id)

    if request.method == "POST" and action == "set_primary_team_access":
        try:
            access_id = int(request.POST.get("access_id") or 0)
        except (TypeError, ValueError):
            access_id = 0

        team_access = get_object_or_404(TeamAccess, pk=access_id, team=team)
        with transaction.atomic():
            TeamAccess.objects.filter(team=team, is_primary=True).update(is_primary=False)
            team_access.is_primary = True
            team_access.is_active = True
            team_access.save(update_fields=["is_primary", "is_active", "updated_at"])
        messages.success(request, f"{team_access.user} is now the primary owner login for {team.name}.")
        return redirect("team_edit", slug=tournament.slug, team_id=team.id)

    # Add retained / pre-auction player to team
    if request.method == "POST" and action == "add_player":
        try:
            player_id = int(request.POST.get("player_id") or 0)
        except (TypeError, ValueError):
            player_id = 0

        price_raw = (request.POST.get("retained_price") or "").strip()
        retained_price = Decimal("0")
        if price_raw:
            try:
                retained_price = Decimal(price_raw)
            except (InvalidOperation, ValueError):
                retained_price = Decimal("0")
        if retained_price < 0:
            retained_price = Decimal("0")

        player = get_object_or_404(Player, pk=player_id, tournament=tournament)

        sold_exists = AuctionLot.objects.filter(
            auction__tournament=tournament,
            player=player,
            status=AuctionLot.Status.SOLD,
        ).exists()
        if sold_exists:
            messages.error(request, "Can't assign: this player is already SOLD in auction.")
            return redirect("team_edit", slug=tournament.slug, team_id=team.id)

        if TeamPlayer.objects.filter(tournament=tournament, player=player).exists():
            messages.error(request, "This player is already assigned to a team.")
            return redirect("team_edit", slug=tournament.slug, team_id=team.id)

        auction = _ensure_auction(tournament)
        lot = AuctionLot.objects.filter(auction=auction, player=player).first()
        if lot and lot.status == AuctionLot.Status.RUNNING:
            messages.error(request, "Can't assign while this player's lot is RUNNING.")
            return redirect("team_edit", slug=tournament.slug, team_id=team.id)

        # Purse check (if configured)
        if (team.purse_total or 0) > 0 and retained_price > 0:
            if (team.purse_remaining or 0) < retained_price:
                messages.error(request, "Team doesn't have enough purse remaining for this retained price.")
                return redirect("team_edit", slug=tournament.slug, team_id=team.id)

        with transaction.atomic():
            TeamPlayer.objects.create(
                tournament=tournament,
                team=team,
                player=player,
                bought_in_auction=None,
                bought_price=retained_price,
                status=TeamPlayer.Status.ACTIVE,
                designation=TeamPlayer.Designation.PLAYER,
            )

            # Deduct purse_remaining if configured
            if (team.purse_total or 0) > 0 and retained_price > 0:
                team.purse_remaining = (team.purse_remaining or 0) - retained_price
                team.save(update_fields=["purse_remaining", "updated_at"])

            # Ensure player won't come up for bidding again.
            if lot and lot.status != AuctionLot.Status.SOLD:
                lot.status = AuctionLot.Status.WITHDRAWN
                lot.save(update_fields=["status", "updated_at"])

        messages.success(request, f"Assigned {player.name} to {team.name}.")
        return redirect("team_edit", slug=tournament.slug, team_id=team.id)

    # Remove retained player (only if not sold via auction)
    if request.method == "POST" and action == "remove_player":
        try:
            tp_id = int(request.POST.get("team_player_id") or 0)
        except (TypeError, ValueError):
            tp_id = 0

        tp = get_object_or_404(TeamPlayer, pk=tp_id, tournament=tournament, team=team)
        if tp.bought_in_auction_id:
            messages.error(request, "Can't remove: this player was bought in auction.")
            return redirect("team_edit", slug=tournament.slug, team_id=team.id)

        refund = Decimal(str(tp.bought_price or 0))
        auction = _ensure_auction(tournament)
        lot = AuctionLot.objects.filter(auction=auction, player=tp.player).first()

        with transaction.atomic():
            tp.delete()

            # Refund purse_remaining if configured
            if (team.purse_total or 0) > 0 and refund > 0:
                team.purse_remaining = (team.purse_remaining or 0) + refund
                team.save(update_fields=["purse_remaining", "updated_at"])

            # Optionally reopen lot if it was withdrawn by retention
            if lot and lot.status == AuctionLot.Status.WITHDRAWN:
                lot.status = AuctionLot.Status.PENDING
                lot.save(update_fields=["status", "updated_at"])

        messages.success(request, f"Removed {tp.player.name} from {team.name}.")
        return redirect("team_edit", slug=tournament.slug, team_id=team.id)

    squad = (
        TeamPlayer.objects.filter(tournament=tournament, team=team)
        .select_related("player")
        .order_by("player__name")
    )

    assigned_ids = TeamPlayer.objects.filter(tournament=tournament).values_list("player_id", flat=True)
    sold_ids = AuctionLot.objects.filter(
        auction__tournament=tournament,
        status=AuctionLot.Status.SOLD,
    ).values_list("player_id", flat=True)

    available_players = (
        Player.objects.filter(tournament=tournament)
        .exclude(id__in=assigned_ids)
        .exclude(id__in=sold_ids)
        .order_by("name")
    )
    access_entries = TeamAccess.objects.filter(team=team).select_related("user").order_by(
        "-is_primary",
        "-is_active",
        "role",
        "user__username",
    )

    return render(
        request,
        "tournaments/team_edit.html",
        {
            "tournament": tournament,
            "team": team,
            "form": form,
            "squad": squad,
            "available_players": available_players,
            "designation_choices": TeamPlayer.Designation.choices,
            "access_form": access_form,
            "access_entries": access_entries,
        },
    )


@login_required
def player_create(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    if request.method == "POST":
        form = PlayerForm(request.POST, request.FILES)
        if form.is_valid():
            player = form.save(commit=False)
            player.tournament = tournament
            player.save()

            if _is_ajax(request):
                player_html = render_to_string(
                    "tournaments/partials/player_card.html",
                    {
                        "p": player,
                        "tournament": tournament,
                        "can_manage": can_manage_tournament(request.user, tournament),
                        "sold_player_ids": set(),
                        "assigned_player_ids": set(),
                    },
                    request=request,
                )
                players_qs = tournament.players.all()
                teams_qs = tournament.teams.all()
                return JsonResponse(
                    {
                        "ok": True,
                        "player": {"id": player.id, "name": player.name},
                        "player_html": player_html,
                        "counts": {
                            "total_teams": teams_qs.count(),
                            "total_players": players_qs.count(),
                            "active_players": players_qs.filter(is_active=True).count(),
                        },
                    }
                )

            messages.success(request, "Player created.")
        else:
            if _is_ajax(request):
                return JsonResponse(
                    {
                        "ok": False,
                        "error": "Could not create player. Please check the inputs.",
                        "errors": form.errors.get_json_data(),
                    },
                    status=400,
                )
            messages.error(request, "Could not create player. Please check the inputs.")

        return redirect("tournament_detail", slug=tournament.slug)

    # Fallback page (not used in modal flow)
    form = PlayerForm()
    return render(
        request,
        "tournaments/player_form.html",
        {"tournament": tournament, "form": form, "mode": "create"},
    )


@login_required
@require_POST
def player_update(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    try:
        player_id = int(request.POST.get("player_id"))
    except (TypeError, ValueError):
        if _is_ajax(request):
            return JsonResponse({"ok": False, "error": "Invalid player."}, status=400)
        messages.error(request, "Invalid player.")
        return redirect("tournament_detail", slug=tournament.slug)

    player = get_object_or_404(Player, pk=player_id, tournament=tournament)
    form = PlayerForm(request.POST, request.FILES, instance=player)
    if form.is_valid():
        form.save()

        if _is_ajax(request):
            sold_player_ids = set(
                AuctionLot.objects.filter(
                    auction__tournament=tournament,
                    player=player,
                    status=AuctionLot.Status.SOLD,
                ).values_list("player_id", flat=True)
            )
            assigned_player_ids = set(
                TeamPlayer.objects.filter(tournament=tournament, player=player).values_list(
                    "player_id", flat=True
                )
            )
            player_html = render_to_string(
                "tournaments/partials/player_card.html",
                {
                    "p": player,
                    "tournament": tournament,
                    "can_manage": can_manage_tournament(request.user, tournament),
                    "sold_player_ids": sold_player_ids,
                    "assigned_player_ids": assigned_player_ids,
                },
                request=request,
            )
            players_qs = tournament.players.all()
            teams_qs = tournament.teams.all()
            return JsonResponse(
                {
                    "ok": True,
                    "player": {"id": player.id, "name": player.name},
                    "player_html": player_html,
                    "counts": {
                        "total_teams": teams_qs.count(),
                        "total_players": players_qs.count(),
                        "active_players": players_qs.filter(is_active=True).count(),
                    },
                }
            )

        messages.success(request, "Player updated.")
    else:
        if _is_ajax(request):
            return JsonResponse(
                {
                    "ok": False,
                    "error": "Could not update player. Please check the inputs.",
                    "errors": form.errors.get_json_data(),
                },
                status=400,
            )
        messages.error(request, "Could not update player. Please check the inputs.")

    return redirect("tournament_detail", slug=tournament.slug)


@login_required
@require_POST
def player_delete(request, slug: str, player_id: int):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    if not can_manage_tournament(request.user, tournament):
        if _is_ajax(request):
            return JsonResponse({"ok": False, "error": "Forbidden"}, status=403)
        messages.error(request, "Forbidden.")
        return redirect("tournament_detail", slug=tournament.slug)

    player = get_object_or_404(Player, pk=player_id, tournament=tournament)

    sold_exists = AuctionLot.objects.filter(
        auction__tournament=tournament,
        player=player,
        status=AuctionLot.Status.SOLD,
    ).exists()
    if sold_exists:
        if _is_ajax(request):
            return JsonResponse(
                {"ok": False, "error": "Can't delete: this player is SOLD in auction."},
                status=400,
            )
        messages.error(request, "Can't delete: this player is SOLD in auction.")
        return redirect("tournament_detail", slug=tournament.slug)

    if TeamPlayer.objects.filter(tournament=tournament, player=player).exists():
        if _is_ajax(request):
            return JsonResponse(
                {"ok": False, "error": "Can't delete: this player is assigned to a team."},
                status=400,
            )
        messages.error(request, "Can't delete: this player is assigned to a team.")
        return redirect("tournament_detail", slug=tournament.slug)

    running_exists = AuctionLot.objects.filter(
        auction__tournament=tournament,
        player=player,
        status=AuctionLot.Status.RUNNING,
    ).exists()
    if running_exists:
        if _is_ajax(request):
            return JsonResponse(
                {"ok": False, "error": "Can't delete: this player's lot is RUNNING."},
                status=400,
            )
        messages.error(request, "Can't delete: this player's lot is RUNNING.")
        return redirect("tournament_detail", slug=tournament.slug)

    player.delete()

    players_qs = tournament.players.all()
    teams_qs = tournament.teams.all()

    if _is_ajax(request):
        return JsonResponse(
            {
                "ok": True,
                "player_id": player_id,
                "counts": {
                    "total_teams": teams_qs.count(),
                    "total_players": players_qs.count(),
                    "active_players": players_qs.filter(is_active=True).count(),
                },
            }
        )

    messages.success(request, "Player deleted.")
    return redirect("tournament_detail", slug=tournament.slug)


@login_required
def player_edit(request, slug: str, player_id: int):
    # Legacy direct-edit page
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)
    player = get_object_or_404(Player, pk=player_id, tournament=tournament)

    if request.method == "POST":
        form = PlayerForm(request.POST, request.FILES, instance=player)
        if form.is_valid():
            form.save()
            messages.success(request, "Player updated.")
            return redirect("tournament_detail", slug=tournament.slug)
    else:
        form = PlayerForm(instance=player)

    return render(
        request,
        "tournaments/player_form.html",
        {"tournament": tournament, "form": form, "player": player, "mode": "edit"},
    )


@login_required
def owner_dashboard(request, slug: str, team_id: int):
    team = get_object_or_404(
        Team.objects.select_related("tournament", "tournament__auction"),
        pk=team_id,
        tournament__slug=slug,
        is_active=True,
    )
    tournament = team.tournament

    if not can_access_team_dashboard(request.user, team):
        messages.error(request, "Forbidden.")
        return redirect("tournament_list")

    owner_access = _owner_access_for_team(request.user, team)
    can_edit_squad = _can_edit_owner_squad(request.user, team, owner_access)
    can_edit_strategy = _can_edit_owner_strategy(request.user, team, owner_access)

    if request.method == "POST" and request.POST.get("action") in {"add_watchlist", "remove_watchlist"}:
        if not can_edit_strategy:
            messages.error(request, "You do not have permission to edit the watchlist.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        action = request.POST.get("action")
        if action == "remove_watchlist":
            try:
                entry_id = int(request.POST.get("watchlist_entry_id") or 0)
            except (TypeError, ValueError):
                entry_id = 0
            entry = get_object_or_404(TeamWatchlist, pk=entry_id, team=team)
            player_name = entry.player.name
            entry.delete()
            messages.success(request, f"Removed {player_name} from watchlist.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        try:
            player_id = int(request.POST.get("player_id") or 0)
        except (TypeError, ValueError):
            player_id = 0

        player = get_object_or_404(Player, pk=player_id, tournament=tournament, is_active=True)
        already_assigned = TeamPlayer.objects.filter(
            tournament=tournament,
            player=player,
        ).exclude(status=TeamPlayer.Status.RELEASED).exists()
        if already_assigned:
            messages.error(request, "This player is already assigned to a team.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        priority = (request.POST.get("priority") or TeamWatchlist.Priority.MEDIUM).strip()
        allowed_priorities = {choice[0] for choice in TeamWatchlist.Priority.choices}
        if priority not in allowed_priorities:
            priority = TeamWatchlist.Priority.MEDIUM

        note = (request.POST.get("note") or "").strip()[:180]
        entry, created = TeamWatchlist.objects.get_or_create(
            team=team,
            player=player,
            defaults={
                "added_by": request.user,
                "priority": priority,
                "note": note,
            },
        )
        if not created:
            entry.priority = priority
            entry.note = note
            entry.added_by = request.user
            entry.save(update_fields=["priority", "note", "added_by", "updated_at"])

        messages.success(request, f"{'Added' if created else 'Updated'} {player.name} on watchlist.")
        return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

    if request.method == "POST" and request.POST.get("action") == "update_owner_squad":
        if not can_edit_squad:
            messages.error(request, "You do not have permission to edit the squad.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        try:
            team_player_id = int(request.POST.get("team_player_id") or 0)
        except (TypeError, ValueError):
            team_player_id = 0

        team_player = get_object_or_404(
            TeamPlayer.objects.select_related("player"),
            pk=team_player_id,
            tournament=tournament,
            team=team,
        )

        designation_raw = (request.POST.get("designation") or "").strip()

        allowed_designations = {choice[0] for choice in TeamPlayer.Designation.choices}
        if not designation_raw:
            designation_raw = TeamPlayer.Designation.PLAYER
        if designation_raw not in allowed_designations:
            messages.error(request, "Invalid designation selected.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        if owner_access and owner_access.linked_player_id == team_player.player_id and designation_raw != team_player.designation:
            messages.error(request, "You cannot change your own designation from the owner dashboard.")
            return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

        updated = False
        with transaction.atomic():
            if team_player.designation != designation_raw:
                if designation_raw == TeamPlayer.Designation.CAPTAIN:
                    TeamPlayer.objects.filter(
                        tournament=tournament,
                        team=team,
                        designation=TeamPlayer.Designation.CAPTAIN,
                    ).exclude(pk=team_player.pk).update(
                        designation=TeamPlayer.Designation.PLAYER,
                        updated_at=timezone.now(),
                    )
                team_player.designation = designation_raw
                team_player.save(update_fields=["designation", "updated_at"])
                updated = True

        if updated:
            messages.success(request, f"Updated {team_player.player.name}.")
        else:
            messages.info(request, "No changes were made.")
        return redirect("owner_dashboard", slug=tournament.slug, team_id=team.id)

    context = _build_owner_dashboard_context(request.user, team)
    return render(request, "tournaments/owner_dashboard.html", context)


@login_required
def teams_dashboard(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    teams = (
        Team.objects.filter(tournament=tournament, is_active=True)
        .order_by("name")
    )

    # Load squad for each team (sold players)
    team_players = (
        TeamPlayer.objects.filter(tournament=tournament)
        .select_related("team", "player")
        .order_by("team__name", "player__name")
    )

    # Group players by team_id
    players_by_team = {}
    for tp in team_players:
        players_by_team.setdefault(tp.team_id, []).append(tp)

    # Compute used = sum(bought_price)
    # (purse_total - purse_remaining is also fine, but sum is more trustworthy)
    spent_by_team = (
        TeamPlayer.objects.filter(tournament=tournament)
        .values("team_id")
        .annotate(spent=Sum("bought_price"), count=Count("id"))
    )
    spent_map = {x["team_id"]: (x["spent"] or 0, x["count"] or 0) for x in spent_by_team}

    enriched = []
    for t in teams:
        spent, cnt = spent_map.get(t.id, (0, 0))
        enriched.append({
            "team": t,
            "spent": spent,
            "players_count": cnt,
            "players": players_by_team.get(t.id, []),
        })

    return render(
        request,
        "tournaments/teams_dashboard.html",
        {"tournament": tournament, "teams": enriched},
    )


@login_required
def export_teams_xlsx(request, slug: str):
    """Export all teams into a single XLSX workbook (one sheet per team)."""

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse(
            "Export requires openpyxl. Install it with: pip install openpyxl",
            status=500,
            content_type="text/plain",
        )

    import re
    from io import BytesIO

    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    teams = list(
        Team.objects.filter(tournament=tournament, is_active=True)
        .order_by("name")
    )

    team_players = (
        TeamPlayer.objects.filter(tournament=tournament)
        .select_related("team", "player")
        .order_by("team__name", "player__name")
    )

    players_by_team = {}
    for tp in team_players:
        players_by_team.setdefault(tp.team_id, []).append(tp)

    wb = Workbook()

    # Remove default sheet
    default_ws = wb.active
    wb.remove(default_ws)

    used_titles = set()

    def safe_sheet_title(name: str) -> str:
        name = (name or "Team").strip() or "Team"
        # Excel sheet title rules: max 31 chars and can't contain: : \\ / ? * [ ]
        name = re.sub(r"[:\\\\/\\?\\*\\[\\]]", " ", name)
        name = re.sub(r"\\s+", " ", name).strip()
        return (name or "Team")[:31]

    for team in teams:
        base_title = safe_sheet_title(team.short_name or team.name)
        title = base_title
        i = 2
        while title in used_titles:
            suffix = f" {i}"
            title = (base_title[: (31 - len(suffix))] + suffix)[:31]
            i += 1
        used_titles.add(title)

        ws = wb.create_sheet(title=title)
        ws.append(["Player Name", "Designation"])

        for tp in players_by_team.get(team.id, []):
            designation = ""
            if tp.designation == TeamPlayer.Designation.OWNER:
                designation = "Owner"
            elif tp.designation == TeamPlayer.Designation.CAPTAIN:
                designation = "Captain"
            ws.append([tp.player.name, designation])

        # Basic formatting (column widths)
        ws.column_dimensions[get_column_letter(1)].width = 32
        ws.column_dimensions[get_column_letter(2)].width = 16
        ws.freeze_panes = "A2"

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"{tournament.slug}_teams.xlsx"
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def export_players_xlsx(request, slug: str):
    """Export all tournament players into a single XLSX sheet."""

    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse(
            "Export requires openpyxl. Install it with: pip install openpyxl",
            status=500,
            content_type="text/plain",
        )

    from io import BytesIO

    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)

    players = list(Player.objects.filter(tournament=tournament).order_by("name"))

    auction = getattr(tournament, "auction", None)

    lot_by_player_id = {}
    if auction:
        lots = (
            AuctionLot.objects.filter(auction=auction)
            .select_related("sold_to_team")
            .order_by("lot_no", "lot_order", "id")
        )
        for lot in lots:
            lot_by_player_id[lot.player_id] = lot

    tp_by_player_id = {
        tp.player_id: tp
        for tp in TeamPlayer.objects.filter(tournament=tournament)
        .select_related("team")
        .order_by("player__name")
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Players"

    headers = [
        "Name",
        # "Role",
        # "City",
        # "Age",
        "Contact",
        # "Active",
        # "Base Price",
        # "Reserve Price",
        # "Team",
        # "Designation",
        # "Lot No",
        # "Lot Status",
        # "Sold To",
        # "Sold Price",
        # "Batting Style",
        # "Bowling Style",
        # "Jersey No",
        # "Jersey Name",
        # "Jersey Size",
        # "Notes",
        # "Stats",
    ]
    ws.append(headers)

    for p in players:
        # lot = lot_by_player_id.get(p.id)
        # tp = tp_by_player_id.get(p.id)
        #
        # stats_val = ""
        # if p.stats:
        #     try:
        #         stats_val = json.dumps(p.stats, ensure_ascii=False)
        #     except Exception:
        #         stats_val = str(p.stats)

        ws.append(
            [
                p.name,
                # p.get_role_display() if p.role else "",
                # p.city or "",
                # p.age or "",
                p.phone or "",
                # "Yes" if p.is_active else "No",
                # p.base_price or 0,
                # p.reserve_price or 0,
                # tp.team.name if tp else "",
                # tp.get_designation_display() if tp else "",
                # lot.lot_no if lot else "",
                # lot.get_status_display() if lot else "",
                # lot.sold_to_team.name if lot and lot.sold_to_team_id else "",
                # lot.sold_price if lot and lot.sold_price is not None else "",
                # p.batting_style or "",
                # p.bowling_style or "",
                # p.jersey_no or "",
                # p.jersey_name or "",
                # p.jersey_size or "",
                # p.notes or "",
                # stats_val,
            ]
        )

    # Basic formatting
    ws.freeze_panes = "A2"
    widths = [
        28,  # Name
        # 14,  # Role
        # 18,  # City
        # 8,   # Age
        # 16,  # Phone
        # 8,   # Active
        # 12,  # Base
        # 12,  # Reserve
        # 20,  # Team
        # 14,  # Designation
        # 8,   # Lot No
        # 14,  # Lot Status
        # 20,  # Sold To
        # 12,  # Sold Price
        # 16,  # Batting
        # 16,  # Bowling
        # 10,  # Jersey No
        # 14,  # Jersey Name
        # 12,  # Jersey Size
        # 26,  # Notes
        # 44,  # Stats
    ]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    out = BytesIO()
    wb.save(out)
    out.seek(0)

    filename = f"{tournament.slug}_players.xlsx"
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def lot_manager(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)
    auction = _ensure_auction(tournament)

    lots = (
        AuctionLot.objects.filter(auction=auction)
        .select_related("player")
        .order_by("lot_no", "lot_order", "id")
    )

    lots_by_no = {}
    for lot in lots:
        lots_by_no.setdefault(lot.lot_no, []).append(lot)

    max_lot_no = max([4] + list(lots_by_no.keys())) if lots_by_no else 4
    lot_numbers = list(range(1, max_lot_no + 1))

    lot_columns = [{"lot_no": n, "lots": lots_by_no.get(n, [])} for n in lot_numbers]

    return render(
        request,
        "tournaments/lot_manager.html",
        {
            "tournament": tournament,
            "auction": auction,
            "lot_columns": lot_columns,
        },
    )


@login_required
@require_POST
def move_lot_player(request, slug: str):
    tournament = get_object_or_404(_filter_tournament_qs(Tournament, request.user), slug=slug)
    auction = _ensure_auction(tournament)

    try:
        payload = json.loads(request.body.decode("utf-8"))
        lot_id = int(payload.get("lot_id"))
        target_lot_no = int(payload.get("target_lot_no"))
    except Exception:
        return JsonResponse({"ok": False, "error": "Invalid payload."}, status=400)

    if target_lot_no < 1:
        return JsonResponse({"ok": False, "error": "Invalid target lot."}, status=400)

    lot = get_object_or_404(AuctionLot, pk=lot_id, auction=auction)
    if lot.status != AuctionLot.Status.PENDING:
        return JsonResponse({"ok": False, "error": "Only pending lots can be moved."}, status=400)

    source_lot_no = lot.lot_no

    with transaction.atomic():
        max_order = (
                AuctionLot.objects.filter(auction=auction, lot_no=target_lot_no)
                .aggregate(Max("lot_order"))
                .get("lot_order__max")
                or 0
        )
        lot.lot_no = target_lot_no
        lot.lot_order = max_order + 1
        lot.save(update_fields=["lot_no", "lot_order", "updated_at"])

        resequence_lot_orders(auction, source_lot_no)
        resequence_lot_orders(auction, target_lot_no)

    return JsonResponse(
        {
            "ok": True,
            "lot_id": lot.id,
            "lot_no": lot.lot_no,
            "lot_order": lot.lot_order,
        }
    )
